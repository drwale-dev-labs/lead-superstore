from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from uuid import UUID
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
from datetime import date

from fastapi import APIRouter, HTTPException, Query, status
from fastapi.responses import StreamingResponse

from app.core.db import get_supabase
from app.schemas.payroll import (
    AddCatchUpRequest,
    PayrollEntryAdjustment,
    PayrollEntryUpdate,
    PayrollPeriodCreate,
    SalaryStructureCreate,
)
from app.services.payroll import (
    already_has_catch_up,
    calendar_days_worked_in_period,
    compute_net_salary,
    fetch_existing_entry_keys,
    fetch_prior_generated_periods,
    find_backdated_periods,
)
from app.services.payroll_entry import StaffPayrollInput, compute_entry_for_staff
from app.services import training_bond
from app.services import payslip as payslip_service
from xhtml2pdf import pisa

router = APIRouter()


# ============================================================================
# Salary structures
# ============================================================================


@router.get("/salary-structures")
def list_salary_structures(staff_id: UUID | None = Query(None)):
    supabase = get_supabase()
    query = supabase.table("salary_structures").select(
        "*, staff!salary_structures_staff_id_fkey(first_name, last_name)"
    )
    if staff_id:
        query = query.eq("staff_id", str(staff_id))
    response = query.order("effective_from", desc=True).execute()
    return {"count": len(response.data), "salary_structures": response.data}


@router.post("/salary-structures", status_code=status.HTTP_201_CREATED)
def create_salary_structure(payload: SalaryStructureCreate):
    """Set a new gross salary for a staff member.

    Closes any open structure by setting its effective_to to the day before
    this new one starts. Also recomputes gross_salary/net_pay on any of this
    staff member's DRAFT payroll entries — otherwise an already-generated
    draft keeps paying out the old salary until someone notices and clicks
    Regenerate. Approved/locked periods are never touched.
    """
    supabase = get_supabase()

    # Close any currently open structure
    open_structures = (
        supabase.table("salary_structures")
        .select("id")
        .eq("staff_id", str(payload.staff_id))
        .is_("effective_to", "null")
        .execute()
    )
    for s in open_structures.data:
        prior = date.fromordinal(payload.effective_from.toordinal() - 1).isoformat()
        supabase.table("salary_structures").update(
            {"effective_to": prior}
        ).eq("id", s["id"]).execute()

    insert_data = payload.model_dump(mode="json")
    response = supabase.table("salary_structures").insert(insert_data).execute()
    structure = response.data[0]

    adjustments = _update_draft_entries_gross_salary(
        supabase, payload.staff_id, payload.gross_salary
    )

    return {**structure, "adjustments": adjustments}


def _update_draft_entries_gross_salary(
    supabase, staff_id, new_gross: Decimal
) -> list[str]:
    """Recompute gross_salary/net_pay on this staff member's draft payroll
    entries after a salary change. Returns notes on what was adjusted.
    """
    draft_entries = (
        supabase.table("payroll_entries")
        .select(
            "id, period_id, working_days, deductions, adjustment_amount, "
            "payroll_periods!inner(status)"
        )
        .eq("staff_id", str(staff_id))
        .eq("payroll_periods.status", "draft")
        .execute()
    )

    notes: list[str] = []
    for entry in draft_entries.data:
        new_net = compute_net_salary(
            gross_salary=Decimal(str(new_gross)),
            working_days=entry["working_days"],
            deductions=Decimal(str(entry["deductions"])),
        )
        new_net += Decimal(str(entry.get("adjustment_amount") or 0))
        supabase.table("payroll_entries").update(
            {"gross_salary": float(new_gross), "net_pay": float(new_net)}
        ).eq("id", entry["id"]).execute()
        _refresh_period_totals(supabase, entry["period_id"])
        notes.append("Updated gross pay and recomputed net pay on a draft payroll entry.")

    return notes


# ============================================================================
# Payroll periods
# ============================================================================


@router.get("/periods")
def list_payroll_periods(
    outlet_id: UUID | None = Query(None),
    status_filter: str | None = Query(None, alias="status"),
):
    supabase = get_supabase()
    query = supabase.table("payroll_periods").select("*, outlets(name)")
    if outlet_id:
        query = query.eq("outlet_id", str(outlet_id))
    if status_filter:
        query = query.eq("status", status_filter)
    response = query.order("period_start", desc=True).execute()
    return {"count": len(response.data), "periods": response.data}


@router.get("/periods/{period_id}")
def get_payroll_period(period_id: UUID):
    """Get a payroll period with all its entries."""
    supabase = get_supabase()

    period = (
        supabase.table("payroll_periods")
        .select("*, outlets(name)")
        .eq("id", str(period_id))
        .single()
        .execute()
    )
    if not period.data:
        raise HTTPException(status_code=404, detail="Payroll period not found")

    entries = (
    supabase.table("payroll_entries")
    .select(
        "*, staff(first_name, last_name, role_id, bank_sort_code, roles(name))"
    )
    .eq("period_id", str(period_id))
    .execute()
)

    return {"period": period.data, "entries": entries.data}


@router.delete("/periods/{period_id}", status_code=status.HTTP_200_OK)
def delete_payroll_period(period_id: UUID):
    """Delete a draft payroll period and its entries.

    Only draft periods can be deleted — approving a period commits loan
    balance decrements and training bond totals, which are not reversible
    here, so an approved/locked period must stay.
    """
    supabase = get_supabase()

    period = (
        supabase.table("payroll_periods")
        .select("status")
        .eq("id", str(period_id))
        .single()
        .execute()
    )
    if not period.data:
        raise HTTPException(status_code=404, detail="Payroll period not found")
    if period.data["status"] != "draft":
        raise HTTPException(
            status_code=400, detail="Only draft periods can be deleted"
        )

    entry_ids = [
        e["id"]
        for e in supabase.table("payroll_entries")
        .select("id")
        .eq("period_id", str(period_id))
        .execute()
        .data
    ]

    # Release advances/fines this period had claimed, same as regenerating
    # does — otherwise they stay stuck as "applied" pointing at a period_id
    # that no longer exists, silently excluded from every future generation.
    supabase.table("salary_advances").update(
        {"status": "pending", "applied_to_period_id": None}
    ).eq("applied_to_period_id", str(period_id)).execute()
    supabase.table("fines").update(
        {"status": "approved", "applied_to_period_id": None}
    ).eq("applied_to_period_id", str(period_id)).execute()

    if entry_ids:
        supabase.table("payroll_entry_deductions").delete().in_(
            "entry_id", entry_ids
        ).execute()
        supabase.table("payroll_entry_bond_items").delete().in_(
            "entry_id", entry_ids
        ).execute()
        supabase.table("payroll_entry_catchups").delete().in_(
            "entry_id", entry_ids
        ).execute()
        supabase.table("payroll_entries").delete().in_("id", entry_ids).execute()

    supabase.table("payroll_periods").delete().eq("id", str(period_id)).execute()

    return {"deleted": True}


@router.post("/periods", status_code=status.HTTP_201_CREATED)
def create_payroll_period(payload: PayrollPeriodCreate):
    """Create a draft payroll period for an outlet."""
    supabase = get_supabase()

    if payload.period_end <= payload.period_start:
        raise HTTPException(
            status_code=400, detail="period_end must be after period_start"
        )

    insert_data = payload.model_dump(mode="json", exclude_none=True)
    response = supabase.table("payroll_periods").insert(insert_data).execute()
    return response.data[0]


@router.post("/periods/{period_id}/generate")
def generate_payroll_entries(period_id: UUID):
    """Auto-create payroll entries for active staff at the outlet.

    For each staff:
      - Pull active salary structure (gross_salary)
      - Snapshot bank details onto the entry
      - Pull pending deductions: active loan installments, pending advances,
        approved fines
      - Sum into entry.deductions, snapshot each into payroll_entry_deductions
      - Compute net = (working_days/30) * gross - deductions
    """
    supabase = get_supabase()

    period = (
        supabase.table("payroll_periods")
        .select("*")
        .eq("id", str(period_id))
        .single()
        .execute()
    )
    if not period.data:
        raise HTTPException(status_code=404, detail="Payroll period not found")
    if period.data["status"] != "draft":
        raise HTTPException(
            status_code=400,
            detail="Cannot regenerate — period is no longer in draft",
        )

    # Atomic lock: flip draft -> generating conditioned on it still being
    # draft. If a second concurrent request races in here, its conditional
    # update matches 0 rows (status is already 'generating') and it's
    # rejected with 409 instead of silently interleaving writes with this
    # run — the original cause of a batch of fine deductions going missing.
    lock_resp = (
        supabase.table("payroll_periods")
        .update({"status": "generating"})
        .eq("id", str(period_id))
        .eq("status", "draft")
        .execute()
    )
    if not lock_resp.data:
        raise HTTPException(
            status_code=409,
            detail="Another generation is already in progress for this period",
        )

    try:
        return _run_generate_payroll_entries(supabase, period_id, period.data)
    finally:
        supabase.table("payroll_periods").update({"status": "draft"}).eq(
            "id", str(period_id)
        ).execute()


def _run_generate_payroll_entries(supabase, period_id: UUID, period_row: dict) -> dict:
    outlet_id = period_row["outlet_id"]
    period_start_str = period_row["period_start"]
    period_end = period_row["period_end"]
    period_start = date.fromisoformat(period_start_str)
    period_end_date = date.fromisoformat(period_end)

    # === Unwind any prior generation ===
    prior_entries = (
        supabase.table("payroll_entries")
        .select("id")
        .eq("period_id", str(period_id))
        .execute()
    )
    if prior_entries.data:
        # Revert advances and fines that pointed at this period
        supabase.table("salary_advances").update(
            {"status": "pending", "applied_to_period_id": None}
        ).eq("applied_to_period_id", str(period_id)).execute()
        supabase.table("fines").update(
            {"status": "approved", "applied_to_period_id": None}
        ).eq("applied_to_period_id", str(period_id)).execute()
        # Cascade-delete entries (and their deduction snapshots)
        supabase.table("payroll_entries").delete().eq(
            "period_id", str(period_id)
        ).execute()

    # === Pull staff eligible for this period ===
    # Active staff (normal case, plus mid-period hires — hired_at trims their days).
    active_staff_resp = (
        supabase.table("staff")
        .select(
            "id, first_name, last_name, bank_name, "
            "bank_account_number, bank_account_name, hired_at, terminated_at"
        )
        .eq("outlet_id", outlet_id)
        .eq("status", "active")
        .execute()
    )
    # Staff terminated during this period still worked part of it and are owed
    # a final prorated paycheck — without this they'd be silently skipped since
    # they're no longer 'active' by the time payroll runs.
    terminated_staff_resp = (
        supabase.table("staff")
        .select(
            "id, first_name, last_name, bank_name, "
            "bank_account_number, bank_account_name, hired_at, terminated_at"
        )
        .eq("outlet_id", outlet_id)
        .eq("status", "terminated")
        .gte("terminated_at", period_start_str)
        .lte("terminated_at", period_end)
        .execute()
    )
    staff_list = active_staff_resp.data + terminated_staff_resp.data
    staff_ids = [s["id"] for s in staff_list]

    # === Batch-fetch everything the per-staff loop needs, in one round trip
    # each, instead of issuing 5-7+ queries per staff (was an N+1 pattern). ===
    structures_resp = (
        supabase.table("salary_structures")
        .select("*")
        .in_("staff_id", staff_ids)
        .lte("effective_from", period_end)
        .order("effective_from", desc=True)
        .execute()
        if staff_ids
        else None
    )
    structures_by_staff: dict[str, list[dict]] = defaultdict(list)
    for s in (structures_resp.data if structures_resp else []):
        structures_by_staff[s["staff_id"]].append(s)

    loans_resp = (
        supabase.table("loans")
        .select("*")
        .in_("staff_id", staff_ids)
        .eq("status", "active")
        .execute()
        if staff_ids
        else None
    )
    loans_by_staff: dict[str, list[dict]] = defaultdict(list)
    for loan in (loans_resp.data if loans_resp else []):
        loans_by_staff[loan["staff_id"]].append(loan)

    advances_resp = (
        supabase.table("salary_advances")
        .select("*")
        .in_("staff_id", staff_ids)
        .eq("status", "pending")
        .execute()
        if staff_ids
        else None
    )
    advances_by_staff: dict[str, list[dict]] = defaultdict(list)
    for adv in (advances_resp.data if advances_resp else []):
        advances_by_staff[adv["staff_id"]].append(adv)

    fines_resp = (
        supabase.table("fines")
        .select("*")
        .in_("staff_id", staff_ids)
        .eq("status", "approved")
        .execute()
        if staff_ids
        else None
    )
    fines_by_staff: dict[str, list[dict]] = defaultdict(list)
    for fine in (fines_resp.data if fines_resp else []):
        fines_by_staff[fine["staff_id"]].append(fine)

    bonds_by_staff = training_bond.fetch_bonds_for_staff(supabase, staff_ids)

    hired_staff_ids = [s["id"] for s in staff_list if s.get("hired_at")]
    earliest_hired_at = (
        min(date.fromisoformat(s["hired_at"]) for s in staff_list if s.get("hired_at"))
        if hired_staff_ids
        else None
    )
    prior_periods = (
        fetch_prior_generated_periods(supabase, outlet_id, earliest_hired_at)
        if earliest_hired_at
        else []
    )
    existing_entry_keys = fetch_existing_entry_keys(
        supabase, [p["id"] for p in prior_periods], hired_staff_ids
    )

    total_gross = Decimal("0")
    total_net = Decimal("0")
    skipped: list[str] = []
    backdated: list[dict] = []

    # Entries/deductions/bond items are computed per staff below (no I/O)
    # then written in bulk after the loop — one insert per table instead of
    # one round trip per staff. At outlet scale (70-100+ staff) the old
    # per-staff insert loop could exceed the frontend's 30s request timeout
    # partway through, leaving some staff (including whoever had a loan/fine
    # queued for later in the list) with no entry at all.
    entries_to_insert: list[dict] = []
    entry_staff_order: list[dict] = []  # parallel to entries_to_insert
    advance_ids_applied: list[str] = []
    fine_ids_applied: list[str] = []

    for staff in staff_list:
        # Find active salary structure
        active = next(
            (
                s
                for s in structures_by_staff.get(staff["id"], [])
                if s["effective_to"] is None or s["effective_to"] >= period_end
            ),
            None,
        )
        if not active:
            skipped.append(
                f"{staff['first_name']} {staff['last_name']} (no salary set)"
            )
            continue

        gross = Decimal(str(active["gross_salary"]))
        hired_at = date.fromisoformat(staff["hired_at"]) if staff.get("hired_at") else None
        terminated_at = (
            date.fromisoformat(staff["terminated_at"]) if staff.get("terminated_at") else None
        )

        # Bond resolution/creation is a genuine side effect (may insert a new
        # training_bonds row) — kept here, outside the pure computation below.
        bond = (
            training_bond.get_or_create_bond(
                supabase, staff["id"], hired_at, existing_bonds=bonds_by_staff
            )
            if hired_at
            else None
        )

        staff_input = StaffPayrollInput(
            id=staff["id"],
            first_name=staff["first_name"],
            last_name=staff["last_name"],
            bank_name=staff.get("bank_name"),
            bank_account_number=staff.get("bank_account_number"),
            bank_account_name=staff.get("bank_account_name"),
            hired_at=hired_at,
            terminated_at=terminated_at,
            gross_salary=gross,
            loans=loans_by_staff.get(staff["id"], []),
            advances=advances_by_staff.get(staff["id"], []),
            fines=fines_by_staff.get(staff["id"], []),
            bond=bond,
        )
        result = compute_entry_for_staff(staff_input, period_start, period_end_date)

        advance_ids_applied.extend(result.advance_ids_to_apply)
        fine_ids_applied.extend(result.fine_ids_to_apply)

        entries_to_insert.append(
            {
                "period_id": str(period_id),
                "staff_id": staff["id"],
                "gross_salary": float(result.gross_salary),
                "working_days": result.working_days,
                "deductions": float(result.total_deductions),
                "net_pay": float(result.net_pay),
                "bank_name": staff.get("bank_name"),
                "bank_account_number": staff.get("bank_account_number"),
                "bank_account_name": staff.get("bank_account_name"),
                "payment_status": "pending",
            }
        )
        entry_staff_order.append(
            {
                "staff": staff,
                "hired_at": hired_at,
                "gross": gross,
                "bond": bond,
                "deduction_items": result.deduction_items,
                "bond_item": result.bond_item,
            }
        )

        total_gross += gross
        total_net += result.net_pay

    # === Bulk-write everything computed above ===
    # Mark advances/fines applied to this period, in as few round trips as
    # the client supports (still N calls if it doesn't support .in_() on
    # update, but at minimum this replaces per-staff serial calls below).
    if advance_ids_applied:
        supabase.table("salary_advances").update(
            {"status": "applied", "applied_to_period_id": str(period_id)}
        ).in_("id", advance_ids_applied).execute()
    if fine_ids_applied:
        supabase.table("fines").update(
            {"status": "applied", "applied_to_period_id": str(period_id)}
        ).in_("id", fine_ids_applied).execute()

    inserted_entries = (
        supabase.table("payroll_entries").insert(entries_to_insert).execute().data
        if entries_to_insert
        else []
    )

    # Supabase preserves insert order in the response, so zip back up with
    # the per-staff data computed above to snapshot deductions/bond items
    # and detect backdated periods against the now-real entry_id.
    all_deduction_items: list[dict] = []
    bond_items_to_record: list[tuple[dict, dict, str]] = []
    for entry_row, staff_data in zip(inserted_entries, entry_staff_order):
        entry_id = entry_row["id"]
        staff = staff_data["staff"]

        for item in staff_data["deduction_items"]:
            item["entry_id"] = entry_id
        all_deduction_items.extend(staff_data["deduction_items"])

        if staff_data["bond_item"] and staff_data["bond"]:
            bond_items_to_record.append(
                (staff_data["bond"], staff_data["bond_item"], entry_id)
            )

        # === Detect backdated periods (Phase 8b) ===
        # Never auto-included — surfaced for HR to review and explicitly approve
        # via POST /entries/{entry_id}/catch-up.
        hired_at = staff_data["hired_at"]
        if hired_at:
            missed = find_backdated_periods(
                prior_periods, staff["id"], hired_at, existing_entry_keys
            )
            for m in missed:
                estimated = compute_net_salary(
                    gross_salary=staff_data["gross"],
                    working_days=m["days_owed"],
                    deductions=Decimal("0"),
                )
                backdated.append(
                    {
                        "staff_id": str(staff["id"]),
                        "staff_name": f"{staff['first_name']} {staff['last_name']}",
                        "entry_id": entry_id,
                        "missed_period_id": m["period_id"],
                        "missed_period_label": (
                            f"{m['period_start'].isoformat()} to {m['period_end'].isoformat()}"
                        ),
                        "days_owed": m["days_owed"],
                        "estimated_amount": float(estimated),
                    }
                )

    if all_deduction_items:
        supabase.table("payroll_entry_deductions").insert(all_deduction_items).execute()

    training_bond.record_bond_items_bulk(
        supabase,
        [(bond["id"], bond_item, entry_id) for bond, bond_item, entry_id in bond_items_to_record],
    )

    supabase.table("payroll_periods").update(
        {"total_gross": float(total_gross), "total_net": float(total_net)}
    ).eq("id", str(period_id)).execute()

    return {
        "period_id": str(period_id),
        "entries_created": len(staff_list) - len(skipped),
        "skipped": skipped,
        "backdated": backdated,
        "total_gross": float(total_gross),
        "total_net": float(total_net),
    }


@router.post("/periods/{period_id}/approve")
def approve_payroll_period(period_id: UUID):
    """Lock the period and commit loan balance decrements."""
    supabase = get_supabase()

    period = (
        supabase.table("payroll_periods")
        .select("status")
        .eq("id", str(period_id))
        .single()
        .execute()
    )
    if not period.data:
        raise HTTPException(status_code=404, detail="Payroll period not found")
    if period.data["status"] != "draft":
        raise HTTPException(status_code=400, detail="Period is not in draft status")

    # Commit loan balance decrements
    loan_deductions = (
        supabase.table("payroll_entry_deductions")
        .select("source_id, amount, payroll_entries!inner(period_id)")
        .eq("source_type", "loan")
        .eq("payroll_entries.period_id", str(period_id))
        .execute()
    )

    per_loan: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for d in loan_deductions.data:
        per_loan[d["source_id"]] += Decimal(str(d["amount"]))

    if per_loan:
        loans_resp = (
            supabase.table("loans")
            .select("id, balance")
            .in_("id", list(per_loan.keys()))
            .execute()
        )
        balances_by_loan = {row["id"]: Decimal(str(row["balance"])) for row in loans_resp.data}
        for loan_id, total in per_loan.items():
            new_balance = balances_by_loan[loan_id] - total
            update = {"balance": float(max(new_balance, Decimal("0")))}
            if new_balance <= 0:
                update["status"] = "paid_off"
            supabase.table("loans").update(update).eq("id", loan_id).execute()

    # Commit training bond running totals
    training_bond.commit_bond_items_for_period(supabase, str(period_id))

    # Lock the period
    response = (
        supabase.table("payroll_periods")
        .update(
            {
                "status": "approved",
                "approved_at": datetime.now(timezone.utc).isoformat(),
            }
        )
        .eq("id", str(period_id))
        .execute()
    )
    return response.data[0]

# ============================================================================
# Payroll entries
# ============================================================================
@router.patch("/entries/{entry_id}")
def update_payroll_entry(entry_id: UUID, payload: PayrollEntryUpdate):
    """Edit gross, working days, or deductions. Net is auto-recomputed."""
    supabase = get_supabase()

    entry = (
        supabase.table("payroll_entries")
        .select("*, payroll_periods(status)")
        .eq("id", str(entry_id))
        .single()
        .execute()
    )
    if not entry.data:
        raise HTTPException(status_code=404, detail="Payroll entry not found")
    if entry.data["payroll_periods"]["status"] != "draft":
        raise HTTPException(
            status_code=400, detail="Cannot edit entries in an approved period"
        )

    gross = Decimal(
        str(
            payload.gross_salary
            if payload.gross_salary is not None
            else entry.data["gross_salary"]
        )
    )
    working_days = (
        payload.working_days
        if payload.working_days is not None
        else entry.data["working_days"]
    )
    deductions = Decimal(
        str(
            payload.deductions
            if payload.deductions is not None
            else entry.data["deductions"]
        )
    )

    net = compute_net_salary(
        gross_salary=gross, working_days=working_days, deductions=deductions
    )
    # Preserve any manual adjustment already set on this entry — otherwise
    # editing gross/days/deductions here would silently wipe out a bonus or
    # withholding someone applied earlier via set_entry_adjustment.
    net += Decimal(str(entry.data.get("adjustment_amount") or 0))

    update_data = {
        "gross_salary": float(gross),
        "working_days": working_days,
        "deductions": float(deductions),
        "net_pay": float(net),
    }
    if payload.notes is not None:
        update_data["notes"] = payload.notes

    response = (
        supabase.table("payroll_entries")
        .update(update_data)
        .eq("id", str(entry_id))
        .execute()
    )

    _refresh_period_totals(supabase, entry.data["period_id"])
    return response.data[0]


@router.patch("/entries/{entry_id}/adjustment")
def set_entry_adjustment(entry_id: UUID, payload: PayrollEntryAdjustment):
    """Set a one-off manual adjustment on a draft entry — positive for a
    bonus/owed back pay, negative to withhold. Recomputes net pay on top of
    the existing gross/working-days/deductions math. Always overwrites any
    prior adjustment on this entry rather than stacking (one adjustment per
    entry per period, with a fresh comment for whatever it currently is).
    """
    supabase = get_supabase()

    entry = (
        supabase.table("payroll_entries")
        .select("*, payroll_periods(status)")
        .eq("id", str(entry_id))
        .single()
        .execute()
    )
    if not entry.data:
        raise HTTPException(status_code=404, detail="Payroll entry not found")
    if entry.data["payroll_periods"]["status"] != "draft":
        raise HTTPException(
            status_code=400, detail="Cannot edit entries in an approved period"
        )

    base_net = compute_net_salary(
        gross_salary=Decimal(str(entry.data["gross_salary"])),
        working_days=entry.data["working_days"],
        deductions=Decimal(str(entry.data["deductions"])),
    )
    new_net = base_net + payload.amount

    response = (
        supabase.table("payroll_entries")
        .update(
            {
                "adjustment_amount": float(payload.amount),
                "adjustment_note": payload.note,
                "net_pay": float(new_net),
            }
        )
        .eq("id", str(entry_id))
        .execute()
    )

    _refresh_period_totals(supabase, entry.data["period_id"])
    return response.data[0]


@router.get("/entries/{entry_id}/payslip")
def download_payslip(entry_id: UUID):
    """Generate a per-employee pay slip PDF for this entry's period.

    Available at any period status — for a draft period the numbers are
    whatever the entry currently holds and can still change before approval.
    """
    supabase = get_supabase()

    entry_resp = (
        supabase.table("payroll_entries")
        .select(
            "gross_salary, working_days, adjustment_amount, adjustment_note, net_pay, "
            "staff(first_name, last_name, roles(name)), "
            "payroll_periods(period_start, period_end, status, outlets(name))"
        )
        .eq("id", str(entry_id))
        .single()
        .execute()
    )
    if not entry_resp.data:
        raise HTTPException(status_code=404, detail="Payroll entry not found")
    entry = entry_resp.data

    deductions_resp = (
        supabase.table("payroll_entry_deductions")
        .select("*")
        .eq("entry_id", str(entry_id))
        .execute()
    )

    staff = entry.get("staff") or {}
    period = entry.get("payroll_periods") or {}
    staff_name = f"{staff.get('first_name', '')} {staff.get('last_name', '')}".strip()
    role_name = (staff.get("roles") or {}).get("name", "—")
    outlet_name = (period.get("outlets") or {}).get("name", "—")

    period_start = date.fromisoformat(period["period_start"])
    period_end = date.fromisoformat(period["period_end"])
    period_label = (
        period_start.strftime("%B %Y")
        if period_start.month == period_end.month
        else f"{period_start.strftime('%d %b %Y')} - {period_end.strftime('%d %b %Y')}"
    )

    html = payslip_service.generate_payslip_html(
        staff_name=staff_name,
        role_name=role_name,
        outlet_name=outlet_name,
        period_label=period_label,
        period_status=period.get("status", ""),
        gross_salary=Decimal(str(entry["gross_salary"])),
        working_days=entry["working_days"],
        deduction_items=deductions_resp.data,
        adjustment_amount=Decimal(str(entry.get("adjustment_amount") or 0)),
        adjustment_note=entry.get("adjustment_note"),
        net_pay=Decimal(str(entry["net_pay"])),
    )

    pdf_buffer = BytesIO()
    pisa_result = pisa.CreatePDF(html, dest=pdf_buffer)
    if pisa_result.err:
        raise HTTPException(status_code=500, detail="Failed to render pay slip to PDF")
    pdf_buffer.seek(0)

    filename = f"{staff_name.replace(' ', '_')}_payslip_{period_end.strftime('%b%Y')}.pdf"

    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/entries/{entry_id}/deductions")
def get_entry_deductions(entry_id: UUID):
    """Show the breakdown of what makes up an entry's total deductions."""
    supabase = get_supabase()
    response = (
        supabase.table("payroll_entry_deductions")
        .select("*")
        .eq("entry_id", str(entry_id))
        .order("source_type")
        .execute()
    )
    return {"count": len(response.data), "items": response.data}


@router.get("/entries/{entry_id}/bond-items")
def get_entry_bond_items(entry_id: UUID):
    """Show training bond deduct/payback items for this entry, if any.

    Deduct items also appear in /deductions (they're part of the deductions
    total); payback items only appear here since they add to net pay.
    """
    supabase = get_supabase()
    response = (
        supabase.table("payroll_entry_bond_items")
        .select("*")
        .eq("entry_id", str(entry_id))
        .execute()
    )
    return {"count": len(response.data), "items": response.data}


@router.get("/entries/{entry_id}/catch-ups")
def get_entry_catch_ups(entry_id: UUID):
    """Show backdated catch-up line items added to this entry, if any."""
    supabase = get_supabase()
    response = (
        supabase.table("payroll_entry_catchups")
        .select("*, payroll_periods!missed_period_id(period_start, period_end)")
        .eq("entry_id", str(entry_id))
        .execute()
    )
    return {"count": len(response.data), "items": response.data}


@router.post("/entries/{entry_id}/catch-up", status_code=status.HTTP_201_CREATED)
def add_catch_up(entry_id: UUID, payload: AddCatchUpRequest):
    """HR explicitly approves a backdated catch-up for one missed period.

    Adds a distinct, traceable line item to THIS (current) entry — the
    original missed period is never reopened or touched. Base pay only; no
    deductions are retroactively computed for the missed period (see
    BUGFIX_RUNBOOK.md Phase 8b).
    """
    supabase = get_supabase()

    entry = (
        supabase.table("payroll_entries")
        .select("*, payroll_periods(status)")
        .eq("id", str(entry_id))
        .single()
        .execute()
    )
    if not entry.data:
        raise HTTPException(status_code=404, detail="Payroll entry not found")
    if entry.data["payroll_periods"]["status"] != "draft":
        raise HTTPException(
            status_code=400, detail="Cannot add a catch-up to an approved period"
        )

    staff_id = entry.data["staff_id"]

    if already_has_catch_up(supabase, staff_id, str(payload.missed_period_id)):
        raise HTTPException(
            status_code=409, detail="A catch-up for this staff/period has already been added"
        )

    missed_period = (
        supabase.table("payroll_periods")
        .select("period_start, period_end")
        .eq("id", str(payload.missed_period_id))
        .single()
        .execute()
    )
    if not missed_period.data:
        raise HTTPException(status_code=404, detail="Missed period not found")

    staff = supabase.table("staff").select("hired_at").eq("id", staff_id).single().execute()
    hired_at = date.fromisoformat(staff.data["hired_at"]) if staff.data.get("hired_at") else None
    if not hired_at:
        raise HTTPException(status_code=400, detail="Staff member has no hire date on file")

    period_start = date.fromisoformat(missed_period.data["period_start"])
    period_end = date.fromisoformat(missed_period.data["period_end"])
    days_owed = calendar_days_worked_in_period(period_start, period_end, hired_at, None)
    if days_owed <= 0:
        raise HTTPException(
            status_code=400, detail="Staff member's hire date does not fall within this period"
        )

    gross = Decimal(str(entry.data["gross_salary"]))
    amount = compute_net_salary(gross_salary=gross, working_days=days_owed, deductions=Decimal("0"))

    supabase.table("payroll_entry_catchups").insert(
        {
            "entry_id": str(entry_id),
            "staff_id": staff_id,
            "missed_period_id": str(payload.missed_period_id),
            "days_owed": days_owed,
            "amount": float(amount),
            "description": (
                f"Backdated catch-up: {period_start.isoformat()} to {period_end.isoformat()}"
            ),
        }
    ).execute()

    new_catch_up_total = Decimal(str(entry.data.get("catch_up_pay") or 0)) + amount
    new_net = Decimal(str(entry.data["net_pay"])) + amount

    updated = (
        supabase.table("payroll_entries")
        .update({"catch_up_pay": float(new_catch_up_total), "net_pay": float(new_net)})
        .eq("id", str(entry_id))
        .execute()
    )

    _refresh_period_totals(supabase, entry.data["period_id"])

    return updated.data[0]


def _refresh_period_totals(supabase, period_id: str) -> None:
    entries = (
        supabase.table("payroll_entries")
        .select("gross_salary, net_pay")
        .eq("period_id", period_id)
        .execute()
    )
    total_gross = sum(Decimal(str(e["gross_salary"])) for e in entries.data)
    total_net = sum(Decimal(str(e["net_pay"])) for e in entries.data)
    supabase.table("payroll_periods").update(
        {"total_gross": float(total_gross), "total_net": float(total_net)}
    ).eq("id", period_id).execute()

# ============================================================================
# Export bank sheet
# ============================================================================
@router.get("/periods/{period_id}/export-review-sheet")
def export_review_sheet(period_id: UUID):
    """Export a human-readable payroll sheet for review — any status, not
    just approved. Meant for sending to someone (e.g. a manager or the
    outlet owner) to check before clicking Approve & lock, unlike
    export-bank-sheet which is the strict bank-upload format for approved
    periods only.
    """
    supabase = get_supabase()

    period_resp = (
        supabase.table("payroll_periods")
        .select("*, outlets(name)")
        .eq("id", str(period_id))
        .single()
        .execute()
    )
    if not period_resp.data:
        raise HTTPException(status_code=404, detail="Payroll period not found")
    period = period_resp.data

    entries_resp = (
        supabase.table("payroll_entries")
        .select(
            "gross_salary, working_days, deductions, net_pay, "
            "staff(first_name, last_name, roles(name))"
        )
        .eq("period_id", str(period_id))
        .execute()
    )
    if not entries_resp.data:
        raise HTTPException(status_code=400, detail="No payroll entries to export for this period")

    outlet_name = (period.get("outlets") or {}).get("name", "Outlet")
    period_start = date.fromisoformat(period["period_start"])
    period_end = date.fromisoformat(period["period_end"])
    period_label = f"{period_start.strftime('%d %b %Y')} - {period_end.strftime('%d %b %Y')}"

    wb = Workbook()
    ws = wb.active
    ws.title = outlet_name[:31]

    ws.append([outlet_name])
    ws.append([f"Payroll period: {period_label}"])
    ws.append([f"Status: {period['status'].upper()}"])
    ws.append([])

    headers = ["Employee", "Role", "Gross Pay", "Days Worked", "Deductions", "Net Pay"]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    total_gross = Decimal("0")
    total_deductions = Decimal("0")
    total_net = Decimal("0")
    for e in sorted(
        entries_resp.data,
        key=lambda e: (
            (e.get("staff") or {}).get("first_name", ""),
            (e.get("staff") or {}).get("last_name", ""),
        ),
    ):
        staff = e.get("staff") or {}
        full_name = f"{staff.get('first_name', '')} {staff.get('last_name', '')}".strip()
        role_name = (staff.get("roles") or {}).get("name", "")
        gross = Decimal(str(e["gross_salary"]))
        deductions = Decimal(str(e["deductions"]))
        net = Decimal(str(e["net_pay"]))
        total_gross += gross
        total_deductions += deductions
        total_net += net
        ws.append([
            full_name,
            role_name,
            float(gross),
            e["working_days"],
            float(deductions),
            float(net),
        ])

    ws.append([])
    ws.append(["Total", "", float(total_gross), "", float(total_deductions), float(total_net)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{outlet_name.replace(' ', '_')}_payroll_review_{period_end.strftime('%b%Y')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/periods/{period_id}/export-bank-sheet")
def export_bank_sheet(period_id: UUID):
    """Export approved payroll entries as a bulk bank-payment xlsx.

    Matches the bank's required column format exactly:
    PaymentAmount | PaymentDate | Reference | Remark | VendorCode |
    VendorName | VendorAcctNumber | VendorBankSortCode
    """
    supabase = get_supabase()

    period_resp = (
        supabase.table("payroll_periods")
        .select("*, outlets(name)")
        .eq("id", str(period_id))
        .single()
        .execute()
    )
    if not period_resp.data:
        raise HTTPException(status_code=404, detail="Payroll period not found")
    period = period_resp.data

    entries_resp = (
        supabase.table("payroll_entries")
        .select(
            "net_pay, bank_account_number, bank_account_name, "
            "staff(first_name, last_name, bank_sort_code)"
        )
        .eq("period_id", str(period_id))
        .execute()
    )
    if not entries_resp.data:
        raise HTTPException(status_code=400, detail="No payroll entries to export for this period")

    period_end = date.fromisoformat(period["period_end"])
    # Default payment date: 5th of the month after the period ends
    if period_end.month == 12:
        pay_date = date(period_end.year + 1, 1, 5)
    else:
        pay_date = date(period_end.year, period_end.month + 1, 5)
    payment_date_str = pay_date.strftime("%d/%B/%Y").upper()

    month_label = period_end.strftime("%B %Y").upper()
    remark = f"SALARY FOR {month_label}"

    outlet_name = (period.get("outlets") or {}).get("name", "OUTLET")
    sheet_name = outlet_name[:31]  # Excel sheet name limit

    missing = [
        e for e in entries_resp.data
        if not e.get("bank_account_number") or not (e.get("staff") or {}).get("bank_sort_code")
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        "PaymentAmount (number format, 2 decimal places (max)",
        "PaymentDate (text format, dd / mmm / yyyy, max. 11 characters)",
        "Reference (optional i.e cells can be left blank, text format, alpha-numeric, max. 20 characters)",
        "Remark (text format, alpha-numeric, max. 25 characters)",
        "VendorCode (text format, max. of 32 characters, e.g staff I.D, RC no. or name)",
        "VendorName (text format, alpha-numeric, max. 50 characters)",
        "VendorAcctNumber (text format, numeric, max. 15 digits)",
        "VendorBankSortCode (text format, 9 digits)",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for e in entries_resp.data:
        staff = e.get("staff") or {}
        full_name = f"{staff.get('first_name', '')} {staff.get('last_name', '')}".strip()
        ws.append([
            round(float(e["net_pay"]), 2),
            payment_date_str,
            "",
            remark[:25],
            full_name[:32],
            (e.get("bank_account_name") or full_name)[:50],
            e.get("bank_account_number") or "",
            staff.get("bank_sort_code") or "",
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{outlet_name.replace(' ', '_')}_salary_{period_end.strftime('%b%Y')}.xlsx"
    headers_out = {"Content-Disposition": f'attachment; filename="{filename}"'}
    if missing:
        headers_out["X-Missing-Bank-Details"] = str(len(missing))

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_out,
    )